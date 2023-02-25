import PySimpleGUI as sg
from fpdf import FPDF
import docx

# Create the template
template = """
MEMORANDUM

TO:      [Acting CD at time of processing]

FROM:   [Acting Director at the time of processing]

SUBJECT: TRANSPORTATION OF MATERIALS ON SHEP PROJECT:  INVOICE IN FAVOUR OF MESSRS [Hauler] ENTERPRISE 

DATE:   [Date of processing]

1.  Please refer to the attached letter dated [date the letter was received] from Messrs [Hauler] on the request for payment for haulage of poles, stay blocks and materials from Tema to project sites in the [Region(s)] Regions.

2.  We confirm that the transporter has satisfactorily executed the work as evidenced by the attached waybills.

3.  Furthermore, review of the submitted invoice indicated that the amount due the transporter under this invoice is GH¢[Invoice amount]. The details are attached Messrs [Hauler] presented [no. of waybills] waybills amounting to GH¢ [waybill amount] for processing. However, waybill numbers [duplicate waybills] were duplicates and not original and as a result were not processed. The revised sum is as a result of the reduction in the number of waybills processed. 

4.  The Power Directorate hereby brings the request for payment of an amount of GH¢[Invoice amount] for the haulage of poles, stay blocks and materials to your attention.

5.  We have attached copies of the relevant documents for your perusal, please.

[Acting Director at the time of processing]
"""

# Create the input form layout
layout = [    [sg.Text("Acting CD at time of processing"), sg.InputText()],
    [sg.Text("Acting Director at the time of processing"), sg.InputText()],
    [sg.Text("Hauler"), sg.InputText()],
    [sg.Text("Date of processing"), sg.InputText()],
    [sg.Text("date the letter was received"), sg.InputText()],
    [sg.Text("Region(s)"), sg.InputText()],
    [sg.Text("Invoice amount"), sg.InputText()],
    [sg.Text("No. of waybills"), sg.InputText()],
    [sg.Text("Waybill amount"), sg.InputText()],
    [sg.Text("Duplicate waybills (comma-separated)"), sg.InputText()],
    [sg.Submit(), sg.Cancel()]
]

# Create the form window
form = sg.Window("Fill in the template").Layout(layout)

# Loop until the form is submitted
while True:
    button, values = form.Read()
    if button == "Submit":
        # Replace the placeholders in the template with the input values
        filled_template = template.replace("[Acting CD at time of processing]", values[0])
        filled_template = filled_template.replace("[Acting Director at the time of processing]", values[1])
        filled_template = filled_template.replace("[Hauler]", values[2])
        filled_template = filled_template.replace("[Date of processing]", values[3])
        filled_template = filled_template.replace("[date the letter was received]", values[4])
        filled_template = filled_template.replace("[Region(s)]", values[5])
        filled_template = filled_template.replace("[Invoice amount]", values[6])
        filled_template = filled_template.replace("[no. of waybills]", values[7])
        filled_template = filled_template.replace("[waybill amount]", values[8])
        filled_template = filled_template.replace("[duplicate waybills]", values[9])

        # Prompt the user for the output format
        output_format = sg.popup_get_text("Enter the desired output format: PDF or DOCX")
        if output_format.lower() == "pdf":
            # Create the PDF object
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("Times", size=12)
            pdf.cell(0, 10, filled_template, 0, 1, align="L")
            pdf.output(f"~/HAULAGE_GENERATED PDFs/{values[2]}.pdf")
            sg.popup("Template filled and exported to PDF.")
        elif output_format.lower() == "docx":
            # Create the .docx file
            doc = docx.Document()
            doc.add_paragraph(filled_template)
            doc.save(f"~/HAULAGE_GENERATED PDFs/{values[2]}.docx")
            sg.popup("Template filled and exported to docx.")
        else:
            sg.popup("Invalid output format. Please enter either PDF or DOCX.")

        break
